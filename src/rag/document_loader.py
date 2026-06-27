from __future__ import annotations

import json as _json
from pathlib import Path

from langchain_community.document_loaders import (
    Docx2txtLoader,
    PyPDFLoader,
    TextLoader,
)
from langchain_core.documents import Document

SUPPORTED_EXTENSIONS = {".docx", ".pdf", ".md", ".txt", ".pptx", ".ipynb", ".py"}


def _load_file(path: Path) -> list[Document]:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        docs = Docx2txtLoader(str(path)).load()
    elif suffix == ".pdf":
        docs = PyPDFLoader(str(path)).load()
    elif suffix == ".pptx":
        docs = _load_pptx(path)
    elif suffix == ".ipynb":
        docs = _load_ipynb(path)
    elif suffix in {".md", ".txt", ".py"}:
        docs = TextLoader(str(path), encoding="utf-8", autodetect_encoding=True).load()
    else:
        return []

    for doc in docs:
        doc.metadata.setdefault("source", path.as_posix())
        doc.metadata.setdefault("title", path.stem)
        doc.metadata.setdefault("doc_type", suffix.lstrip("."))
    return docs


def _load_pptx(path: Path) -> list[Document]:
    """PPTX loader: 提取每页 slide 的文本 shape，跳过空 slide。"""
    try:
        from pptx import Presentation  # noqa: F811 - optional dependency
    except ImportError:
        return []
    docs: list[Document] = []
    try:
        prs = Presentation(str(path))
        for slide_idx, slide in enumerate(prs.slides, start=1):
            texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        t = para.text.strip()
                        if t:
                            texts.append(t)
            full_text = "\n".join(texts).strip()
            if not full_text:
                continue
            docs.append(Document(
                page_content=full_text,
                metadata={
                    "slide_number": slide_idx,
                    "source": path.as_posix(),
                    "title": path.stem,
                    "doc_type": "pptx",
                },
            ))
    except Exception:
        pass
    return docs


def _load_ipynb(path: Path) -> list[Document]:
    """IPYNB loader: 提取 markdown + code cell，丢弃 outputs。"""
    docs: list[Document] = []
    try:
        nb = _json.loads(path.read_text(encoding="utf-8"))
        for cell_idx, cell in enumerate(nb.get("cells", [])):
            ctype = cell.get("cell_type", "")
            source = "".join(cell.get("source", []))
            if not source.strip():
                continue
            if ctype not in ("markdown", "code"):
                continue
            docs.append(Document(
                page_content=source,
                metadata={
                    "notebook_cell_index": cell_idx,
                    "cell_type": ctype,
                    "source": path.as_posix(),
                    "title": path.stem,
                    "doc_type": "ipynb",
                },
            ))
    except Exception:
        pass
    return docs


def load_documents(data_dir: str | Path) -> list[Document]:
    root = Path(data_dir)
    if not root.exists():
        raise FileNotFoundError(f"Data directory does not exist: {root}")
    if not root.is_dir():
        raise NotADirectoryError(f"Data path is not a directory: {root}")

    documents: list[Document] = []
    for path in sorted(root.rglob("*")):
        if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS:
            documents.extend(_load_file(path))
    return documents
